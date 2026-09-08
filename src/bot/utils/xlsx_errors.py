"""Generate XLSX report with validation/upload errors."""

from __future__ import annotations

import io
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


@dataclass
class ProductError:
    """Ошибка обработки одного товара."""

    row: int
    sku: str | None
    error: str


def generate_error_report(errors: list[ProductError]) -> bytes:
    """Генерирует XLSX с колонками: Строка, SKU, Ошибка."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ошибки"

    headers = ["Строка", "SKU", "Ошибка"]
    header_font = Font(bold=True, color="FFFFFF")
    error_fill = PatternFill("solid", fgColor="FFFEEEEE")

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = PatternFill("solid", fgColor="FFCC0000")

    for row_idx, err in enumerate(errors, start=2):
        ws.cell(row=row_idx, column=1, value=err.row)
        ws.cell(row=row_idx, column=2, value=err.sku or "")
        cell = ws.cell(row=row_idx, column=3, value=err.error)
        cell.fill = error_fill

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
