"""Parse uploaded XLSX files into validated Pydantic models."""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from bot.core.exceptions import XLSXFormatError, XLSXValidationError
from bot.utils.schemas import PriceUploadRow, ProductUploadRow, StockUploadRow

PRODUCT_HEADER_TO_FIELD: dict[str, str] = {
    "sku (артикул)": "sku",
    "название": "name",
    "category id": "category_id",
    "цена (₽)": "price",
    "старая цена (₽)": "old_price",
    "штрих-код": "barcode",
    "вес (г)": "weight",
    "описание": "description",
}

STOCKS_HEADER_TO_FIELD: dict[str, str] = {
    "sku (артикул)": "sku",
    "warehouse id": "warehouse_id",
    "остаток (шт)": "stock",
}

PRICES_HEADER_TO_FIELD: dict[str, str] = {
    "sku (артикул)": "sku",
    "цена (₽)": "price",
    "старая цена (₽)": "old_price",
}


def _parse_cell(value: object, field: str) -> object:
    if value is None or value == "":
        return None

    if field in ("category_id", "weight", "stock"):
        try:
            return int(value)
        except (ValueError, TypeError) as e:
            raise XLSXValidationError(
                f"Ожидается целое число, получено {value!r}", column=field
            ) from e

    if field in ("price", "old_price"):
        try:
            return Decimal(str(value).replace(",", ".").replace(" ", ""))
        except Exception as e:
            raise XLSXValidationError(f"Ожидается число, получено {value!r}", column=field) from e

    return str(value).strip()


def _parse_sheet(
    data: bytes,
    sheet_name: str,
    model: type,
    header_to_field: dict[str, str],
) -> list:
    """Общий парсер: открывает файл, ищет лист, парсит строки через model."""
    try:
        wb = load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as e:
        raise XLSXFormatError(f"Не удалось открыть файл: {e}") from e

    if sheet_name not in wb.sheetnames:
        raise XLSXFormatError(f"Лист «{sheet_name}» не найден в файле")

    ws = wb[sheet_name]

    headers_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if headers_row is None:
        raise XLSXFormatError("Файл пустой")

    headers = [str(h).strip().lower() if h else "" for h in headers_row]
    field_map: list[str | None] = [header_to_field.get(h) for h in headers]
    if all(f is None for f in field_map):
        raise XLSXFormatError("Заголовки не распознаны. Скачай свежий шаблон.")

    rows: list = []
    for row_idx, raw in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None or v == "" for v in raw):
            continue

        row_data: dict = {}
        for col_idx, value in enumerate(raw):
            field = field_map[col_idx] if col_idx < len(field_map) else None
            if field is None:
                continue
            try:
                row_data[field] = _parse_cell(value, field)
            except XLSXValidationError as e:
                e.row = row_idx
                raise

        try:
            row = model(**row_data)
        except Exception as e:
            raise XLSXValidationError(f"Ошибка валидации: {e}", row=row_idx) from e
        rows.append(row)

    return rows


def parse_product_upload(data: bytes) -> list[ProductUploadRow]:
    return _parse_sheet(data, "Товары", ProductUploadRow, PRODUCT_HEADER_TO_FIELD)


def parse_stocks_upload(data: bytes) -> list[StockUploadRow]:
    return _parse_sheet(data, "Остатки", StockUploadRow, STOCKS_HEADER_TO_FIELD)


def parse_prices_upload(data: bytes) -> list[PriceUploadRow]:
    return _parse_sheet(data, "Цены", PriceUploadRow, PRICES_HEADER_TO_FIELD)
