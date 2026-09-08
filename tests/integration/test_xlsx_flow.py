"""Integration test: template → parse → save as Product drafts."""

from io import BytesIO

import pytest
from cryptography.fernet import Fernet
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from bot.db.base import Base
from bot.db.models import *  # noqa: F401,F403  # register all models
from bot.db.models import Product
from bot.repositories.product import ProductRepository
from bot.repositories.seller import SellerRepository
from bot.repositories.user import UserRepository
from bot.services.encryption import EncryptionService
from bot.utils.xlsx_parser import parse_product_upload
from bot.utils.xlsx_template import generate_product_template


@pytest.fixture
def enc() -> EncryptionService:
    return EncryptionService(Fernet.generate_key().decode())


@pytest.fixture
async def session() -> AsyncSession:
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(bind=engine, expire_on_commit=False)
    async with factory() as s:
        yield s
    await engine.dispose()


def _make_xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    headers = ["SKU (артикул)", "Название", "Category ID", "Цена (₽)"]
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def test_parse_to_pydantic():
    """Сгенерированный шаблон → ProductUploadRow."""
    data = generate_product_template()
    rows = parse_product_upload(data)
    assert len(rows) == 1
    assert rows[0].sku == "SKU-001"


async def test_persist_parsed_rows(session: AsyncSession, enc: EncryptionService):
    """Parsed rows сохраняются в БД как Product draft."""
    user_repo = UserRepository(session)
    user, _ = await user_repo.get_or_create(telegram_id=1)
    await session.commit()

    seller_repo = SellerRepository(session, encryption=enc)
    seller = await seller_repo.create(user.id, "S", "C", "K")
    await session.commit()

    data = _make_xlsx(
        [
            ["SKU-A", "Товар A", 100, 500],
            ["SKU-B", "Товар B", 200, 1000],
            ["SKU-C", "Товар C", 300, 1500],
        ]
    )
    rows = parse_product_upload(data)
    assert len(rows) == 3

    prod_repo = ProductRepository(session)
    for r in rows:
        await prod_repo.add(
            Product(
                seller_id=seller.id,
                offer_id=r.sku,
                sku=0,
                name=r.name,
                category_id=r.category_id,
            )
        )
    await session.commit()

    saved = await prod_repo.list_for_seller(seller.id)
    assert {p.offer_id for p in saved} == {"SKU-A", "SKU-B", "SKU-C"}


async def test_template_to_persistence_full_flow(session: AsyncSession, enc: EncryptionService):
    """Полный flow: шаблон → парсинг → сохранение в БД."""
    user_repo = UserRepository(session)
    user, _ = await user_repo.get_or_create(telegram_id=2)
    await session.commit()

    seller_repo = SellerRepository(session, encryption=enc)
    seller = await seller_repo.create(user.id, "Shop", "CID", "KEY")
    await session.commit()

    # Берём шаблон и добавляем к нему 2 строки данных
    template_data = generate_product_template()
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(template_data))
    ws = wb["Товары"]
    ws.append(["SKU-100", "Карандаш", 111, 50])
    ws.append(["SKU-200", "Ручка", 222, 120])
    buf = BytesIO()
    wb.save(buf)

    rows = parse_product_upload(buf.getvalue())
    assert len(rows) == 3  # 1 example + 2 added

    prod_repo = ProductRepository(session)
    for r in rows:
        existing = await prod_repo.get_by_offer_id(seller.id, r.sku)
        if existing is None:
            await prod_repo.add(
                Product(
                    seller_id=seller.id,
                    offer_id=r.sku,
                    sku=0,
                    name=r.name,
                    category_id=r.category_id,
                )
            )
    await session.commit()

    all_p = await prod_repo.list_for_seller(seller.id)
    assert {p.offer_id for p in all_p} == {"SKU-001", "SKU-100", "SKU-200"}
