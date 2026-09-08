"""Tests for stocks/prices catalog pagination."""

from decimal import Decimal
from io import BytesIO
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

from bot.bot_app.handlers.prices import format_prices_page
from bot.bot_app.handlers.stocks import format_stocks_page
from bot.ozon.client import OzonAPIClient, OzonCredentials
from bot.ozon.modules.prices import PricesModule
from bot.ozon.modules.stocks import StocksModule
from bot.ozon.schemas.prices import PriceInfo
from bot.ozon.schemas.stocks import StockItem
from bot.utils.xlsx_template import build_prices_catalog_xlsx, build_stocks_catalog_xlsx


def test_format_stocks_page_groups_warehouses():
    items = [
        StockItem(product_id=1, offer_id="A<B>", warehouse_name="fbs", present=3),
        StockItem(product_id=1, offer_id="A<B>", warehouse_name="fbo", present=2),
    ]
    text = format_stocks_page("каб1", items, page=0, total=3201, page_size=25)
    assert "1–1" in text
    assert "3201" in text
    assert "<code>A&lt;B&gt;</code> — 5 шт." in text
    assert "fbs: 3" in text


def test_format_prices_page_escapes_and_shows_old():
    items = [PriceInfo(offer_id="X&Y", price=Decimal("100"), old_price=Decimal("120"))]
    text = format_prices_page("каб1", items, page=0, total=50, page_size=25)
    assert "<code>X&amp;Y</code> — 100 ₽ (старая: 120)" in text
    assert "1–1" in text


async def test_stocks_iter_all_follows_cursor():
    module = StocksModule(OzonAPIClient(OzonCredentials("1", "k")))
    pages = [
        {
            "items": [
                {"product_id": 1, "offer_id": "A", "stocks": [{"type": "fbs", "present": 1}]}
            ],
            "total": 2,
            "cursor": "c1",
        },
        {
            "items": [
                {"product_id": 2, "offer_id": "B", "stocks": [{"type": "fbs", "present": 4}]}
            ],
            "total": 2,
            "cursor": "",
        },
    ]
    with patch.object(module._client, "request", AsyncMock(side_effect=pages)):
        catalog = await module.iter_all(page_size=1)
    assert [i.offer_id for i in catalog.result] == ["A", "B"]
    assert catalog.has_next is False


async def test_prices_iter_all_follows_cursor():
    module = PricesModule(OzonAPIClient(OzonCredentials("1", "k")))
    pages = [
        {
            "items": [{"offer_id": "A", "product_id": 1, "price": {"price": "10"}}],
            "total": 2,
            "cursor": "c1",
        },
        {
            "items": [{"offer_id": "B", "product_id": 2, "price": {"price": "20"}}],
            "total": 2,
            "cursor": "",
        },
    ]
    with patch.object(module._client, "request", AsyncMock(side_effect=pages)):
        catalog = await module.iter_all(page_size=1)
    assert [i.offer_id for i in catalog.result] == ["A", "B"]
    assert catalog.result[0].price == Decimal("10")


def test_stocks_catalog_xlsx_rows():
    data = build_stocks_catalog_xlsx(
        [StockItem(product_id=1, offer_id="SKU-1", warehouse_name="fbs", present=8, reserved=1)]
    )
    ws = load_workbook(BytesIO(data))["Остатки"]
    assert ws.cell(row=2, column=1).value == "SKU-1"
    assert ws.cell(row=2, column=4).value == 8


def test_prices_catalog_xlsx_rows():
    data = build_prices_catalog_xlsx(
        [PriceInfo(offer_id="SKU-1", product_id=9, price=Decimal("99.50"))]
    )
    ws = load_workbook(BytesIO(data))["Цены"]
    assert ws.cell(row=2, column=1).value == "SKU-1"
    assert ws.cell(row=2, column=3).value == "99.50"
