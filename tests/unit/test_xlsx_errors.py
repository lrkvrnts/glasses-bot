"""Tests for XLSX error report generator."""

from io import BytesIO

from openpyxl import load_workbook

from bot.utils.xlsx_errors import ProductError, generate_error_report


def test_generate_error_report_returns_xlsx():
    errors = [
        ProductError(row=2, sku="SKU-1", error="Empty price"),
        ProductError(row=5, sku=None, error="Invalid SKU"),
    ]
    data = generate_error_report(errors)
    assert isinstance(data, bytes)
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Строка"
    assert ws.cell(row=2, column=1).value == 2
    assert ws.cell(row=3, column=1).value == 5


def test_generate_empty_errors_returns_headers():
    data = generate_error_report([])
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Строка"
    # no data rows
    assert ws.cell(row=2, column=1).value is None


def test_generate_error_report_preserves_sku():
    errors = [ProductError(row=3, sku="ABC-123", error="Bad category")]
    data = generate_error_report(errors)
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    assert ws.cell(row=2, column=2).value == "ABC-123"
    assert ws.cell(row=2, column=3).value == "Bad category"


def test_generate_error_report_handles_empty_sku():
    errors = [ProductError(row=4, sku=None, error="Validation failed")]
    data = generate_error_report(errors)
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    # None → пустое значение в ячейке
    assert ws.cell(row=2, column=2).value in (None, "")
