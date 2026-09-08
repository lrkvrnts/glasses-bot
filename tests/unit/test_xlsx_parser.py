"""Tests for XLSX parser."""

from io import BytesIO

import pytest
from openpyxl import Workbook

from bot.core.exceptions import XLSXFormatError, XLSXValidationError
from bot.utils.xlsx_parser import parse_product_upload
from bot.utils.xlsx_template import generate_product_template


def test_parse_template_example_row():
    """Парсер читает пример из шаблона."""
    data = generate_product_template()
    rows = parse_product_upload(data)
    assert len(rows) == 1
    assert rows[0].sku == "SKU-001"
    assert rows[0].name == "Пример товара"
    assert rows[0].price == 1500
    assert rows[0].old_price == 2000


def test_parse_invalid_xlsx_raises():
    """Битый файл → XLSXFormatError."""
    with pytest.raises(XLSXFormatError):
        parse_product_upload(b"not-an-xlsx")


def test_parse_missing_required_field_raises():
    """Без обязательного поля → XLSXValidationError."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    headers = ["SKU (артикул)", "Название", "Category ID", "Цена (₽)"]
    ws.append(headers)
    ws.append(["", "Test", 123, 1000])  # пустой SKU

    buf = BytesIO()
    wb.save(buf)

    with pytest.raises(XLSXValidationError) as exc_info:
        parse_product_upload(buf.getvalue())
    assert exc_info.value.row == 2


def test_parse_negative_price_raises():
    """Отрицательная цена → XLSXValidationError."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    ws.append(["SKU (артикул)", "Название", "Category ID", "Цена (₽)"])
    ws.append(["SKU-1", "Test", 123, -100])

    buf = BytesIO()
    wb.save(buf)

    with pytest.raises(XLSXValidationError):
        parse_product_upload(buf.getvalue())


def test_parse_empty_rows_skipped():
    """Полностью пустые строки пропускаются."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    ws.append(["SKU (артикул)", "Название", "Category ID", "Цена (₽)"])
    ws.append(["SKU-1", "Test", 123, 100])
    ws.append([None, None, None, None])
    ws.append(["SKU-2", "Test 2", 124, 200])

    buf = BytesIO()
    wb.save(buf)

    rows = parse_product_upload(buf.getvalue())
    assert len(rows) == 2


def test_parse_multiple_rows():
    """Парсинг 3 валидных строк."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    ws.append(["SKU (артикул)", "Название", "Category ID", "Цена (₽)"])
    ws.append(["SKU-1", "A", 1, 100])
    ws.append(["SKU-2", "B", 2, 200])
    ws.append(["SKU-3", "C", 3, 300])

    buf = BytesIO()
    wb.save(buf)

    rows = parse_product_upload(buf.getvalue())
    assert [r.sku for r in rows] == ["SKU-1", "SKU-2", "SKU-3"]


def test_parse_missing_sheet_raises():
    """Файл без листа 'Товары' → XLSXFormatError."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Other"
    ws.append(["a", "b"])
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(XLSXFormatError):
        parse_product_upload(buf.getvalue())


def test_parse_unknown_headers_raises():
    """Неизвестные заголовки → XLSXFormatError."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    ws.append(["Foo", "Bar"])
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(XLSXFormatError):
        parse_product_upload(buf.getvalue())


def test_parse_comma_decimal():
    """Цена с запятой как десятичный разделитель корректно парсится."""
    from decimal import Decimal

    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    ws.append(["SKU (артикул)", "Название", "Category ID", "Цена (₽)"])
    ws.append(["SKU-1", "Test", 1, "99,99"])
    buf = BytesIO()
    wb.save(buf)
    rows = parse_product_upload(buf.getvalue())
    assert rows[0].price == Decimal("99.99")


def test_parse_unicode_sku_and_name():
    """Кириллица в SKU/названии корректно парсится."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    ws.append(["SKU (артикул)", "Название", "Category ID", "Цена (₽)"])
    ws.append(["АРТ-001", "Карандаш чернографитный", 100, 50])
    buf = BytesIO()
    wb.save(buf)
    rows = parse_product_upload(buf.getvalue())
    assert rows[0].sku == "АРТ-001"
    assert "Карандаш" in rows[0].name


def test_parse_zero_price_raises():
    """Нулевая цена → XLSXValidationError (gt=0)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    ws.append(["SKU (артикул)", "Название", "Category ID", "Цена (₽)"])
    ws.append(["SKU-1", "Test", 100, 0])
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(XLSXValidationError):
        parse_product_upload(buf.getvalue())


def test_parse_string_category_id_raises():
    """Нечисловой Category ID → XLSXValidationError."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    ws.append(["SKU (артикул)", "Название", "Category ID", "Цена (₽)"])
    ws.append(["SKU-1", "Test", "abc", 100])
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(XLSXValidationError):
        parse_product_upload(buf.getvalue())


def test_parse_stocks_upload_valid_rows():
    from bot.utils.xlsx_parser import parse_stocks_upload

    wb = Workbook()
    ws = wb.active
    ws.title = "Остатки"
    ws.append(["SKU (артикул)", "Warehouse ID", "Остаток (шт)"])
    ws.append(["SKU-1", "wh-1", 100])
    ws.append(["SKU-2", "wh-1", 0])
    buf = BytesIO()
    wb.save(buf)

    rows = parse_stocks_upload(buf.getvalue())
    assert len(rows) == 2
    assert rows[0].sku == "SKU-1"
    assert rows[0].warehouse_id == "wh-1"
    assert rows[0].stock == 100
    assert rows[1].stock == 0


def test_parse_stocks_negative_raises():
    from bot.utils.xlsx_parser import parse_stocks_upload

    wb = Workbook()
    ws = wb.active
    ws.title = "Остатки"
    ws.append(["SKU (артикул)", "Warehouse ID", "Остаток (шт)"])
    ws.append(["SKU-1", "wh-1", -5])
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(XLSXValidationError):
        parse_stocks_upload(buf.getvalue())


def test_parse_prices_upload_valid_rows():
    from decimal import Decimal

    from bot.utils.xlsx_parser import parse_prices_upload

    wb = Workbook()
    ws = wb.active
    ws.title = "Цены"
    ws.append(["SKU (артикул)", "Цена (₽)", "Старая цена (₽)"])
    ws.append(["SKU-1", 1500, 2000])
    ws.append(["SKU-2", 500, None])
    buf = BytesIO()
    wb.save(buf)

    rows = parse_prices_upload(buf.getvalue())
    assert len(rows) == 2
    assert rows[0].price == Decimal("1500")
    assert rows[0].old_price == Decimal("2000")
    assert rows[1].old_price is None


def test_parse_prices_missing_price_raises():
    from bot.utils.xlsx_parser import parse_prices_upload

    wb = Workbook()
    ws = wb.active
    ws.title = "Цены"
    ws.append(["SKU (артикул)", "Цена (₽)", "Старая цена (₽)"])
    ws.append(["SKU-1", None, 2000])
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(XLSXValidationError):
        parse_prices_upload(buf.getvalue())
