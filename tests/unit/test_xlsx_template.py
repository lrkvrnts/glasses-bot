"""Tests for XLSX template generator."""

from io import BytesIO

from openpyxl import load_workbook

from bot.utils.xlsx_template import (
    PRODUCT_TEMPLATE_COLUMNS,
    generate_product_template,
    get_template_filename,
)


def test_generate_product_template_returns_bytes():
    data = generate_product_template()
    assert isinstance(data, bytes)
    assert len(data) > 0


def test_template_opens_as_xlsx():
    data = generate_product_template()
    wb = load_workbook(BytesIO(data))
    assert "Товары" in wb.sheetnames
    assert "Инструкция" in wb.sheetnames


def test_template_has_all_columns():
    data = generate_product_template()
    wb = load_workbook(BytesIO(data))
    ws = wb["Товары"]
    headers = [cell.value for cell in ws[1]]
    expected = [c["title"] for c in PRODUCT_TEMPLATE_COLUMNS]
    assert headers == expected


def test_template_has_example_row():
    data = generate_product_template()
    wb = load_workbook(BytesIO(data))
    ws = wb["Товары"]
    assert ws.cell(row=2, column=1).value == "SKU-001"


def test_template_filename():
    assert get_template_filename() == "products_template.xlsx"


def test_template_instructions_sheet_has_title():
    data = generate_product_template()
    wb = load_workbook(BytesIO(data))
    ws = wb["Инструкция"]
    assert "Как заполнять" in str(ws["A1"].value)


def test_generate_stocks_template():
    from bot.utils.xlsx_template import (
        generate_stocks_template,
        get_stocks_template_filename,
    )

    data = generate_stocks_template()
    assert isinstance(data, bytes)
    wb = load_workbook(BytesIO(data))
    assert "Остатки" in wb.sheetnames
    ws = wb["Остатки"]
    headers = [cell.value for cell in ws[1]]
    assert "SKU (артикул)" in headers
    assert "Warehouse ID" in headers
    assert "Остаток (шт)" in headers
    assert ws.cell(row=2, column=1).value == "SKU-001"
    assert get_stocks_template_filename() == "stocks_template.xlsx"


def test_generate_prices_template():
    from bot.utils.xlsx_template import (
        generate_prices_template,
        get_prices_template_filename,
    )

    data = generate_prices_template()
    assert isinstance(data, bytes)
    wb = load_workbook(BytesIO(data))
    assert "Цены" in wb.sheetnames
    ws = wb["Цены"]
    headers = [cell.value for cell in ws[1]]
    assert "SKU (артикул)" in headers
    assert "Цена (₽)" in headers
    assert ws.cell(row=2, column=2).value == "1500"
    assert get_prices_template_filename() == "prices_template.xlsx"
